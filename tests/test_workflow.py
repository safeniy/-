"""Functional tests use generated images and an explicitly substituted test provider, never fake production OCR."""
import copy
import io
import json
import time
import zipfile

import fitz
import pytest
from docx import Document
from fastapi.testclient import TestClient
from lxml import etree
from PIL import Image, ImageDraw

from notebook.app import create_app
from notebook.diagrams import render_svg, template
from notebook.exporting import MATH_NS, formula_element
from notebook.provider import Provider, ProviderError
from notebook.store import now, uid

FORMULAS=[r'\langle v_x\rangle=\frac{\Delta x}{\Delta t}',r'v_x=\lim_{\Delta t\to0}\frac{\Delta x}{\Delta t}=\frac{dx}{dt}',r'x(t)=x_0+\int_{t_0}^{t}v_x\,dt',r'\vec a+\vec b=\vec c',r'\vec a\cdot\vec b=|\vec a|\,|\vec b|\cos\alpha',r'\vec d=\vec a\times\vec b,\qquad |\vec d|=|\vec a|\,|\vec b|\sin\alpha',r'\vec e_a=\frac{\vec a}{|\vec a|}']


def photo(seed=0,format='PNG'):
    img=Image.new('RGB',(700,1000),'white'); draw=ImageDraw.Draw(img)
    for x in range(0,700,25): draw.line((x,0,x,1000),fill='#cddce1')
    for y in range(0,1000,25): draw.line((0,y,700,y),fill='#cddce1')
    draw.line((100+seed*50,200,500,300+seed*70),fill='black',width=5)
    out=io.BytesIO();img.save(out,format=format);return out.getvalue()


def block(ident='b1',text='тем повышается',typ='text',**kw):
    return {'id':ident,'type':typ,'text':text,'original_text':text,'latex':'','bbox':[.1,.1,.8,.2],**kw}


def document(blocks=None,question=True):
    return {'blocks':blocks or [block()],'questions':[{'id':'q1','block_id':'b1','kind':'uncertain','prompt':'Тем — это температура?','options':['температура'],'original':'тем','bbox':[.1,.1,.15,.05],'resolved':False,'answer':''}] if question else [],'notes':''}


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.delenv('OCR_API_KEY',raising=False);monkeypatch.delenv('APP_PASSWORD',raising=False)
    app=create_app(tmp_path/'data')
    with TestClient(app) as c: yield c


def project(c):
    r=c.post('/api/projects',json={'title':'Физика','subject':'Физика'});assert r.status_code==201,r.text
    return r.json()['id']


def upload(c,pid,data=None,name='page.png'):
    r=c.post(f'/api/projects/{pid}/upload',files={'files':(name,data or photo(),'image/png')});assert r.status_code==200,r.text
    return r.json()['pages'][0]


def edit(c,p,doc):
    r=c.put(f'/api/pages/{p["id"]}',json={'version':p['version'],'document':doc,'source_number':p['source_number']});assert r.status_code==200,r.text
    return r.json()


def respond(c,p,message,**kw):
    r=c.post(f'/api/pages/{p["id"]}/reply',json={'version':p['version'],'message':message,**kw});assert r.status_code==200,r.text
    return r.json()


def test_manual_review_persistence_and_word(tmp_path,monkeypatch):
    monkeypatch.delenv('OCR_API_KEY',raising=False)
    root=tmp_path/'persistent'
    with TestClient(create_app(root)) as c:
        pid=project(c); first=edit(c,upload(c,pid),document()); second=edit(c,upload(c,pid,photo(1)),document())
        result=respond(c,second,'Тем — это температура',question_id='q1')
        second=result['page'];assert second['document']['blocks'][0]['text']=='температура повышается'
        assert second['document']['blocks'][0]['original_text']=='тем повышается'
        second=respond(c,second,'Всё верно')['page'];assert second['status']=='confirmed'
        old=c.get(f'/api/pages/{first["id"]}').json();assert old['status']=='review' and not old['document']['questions'][0]['resolved']
        assert c.post(f'/api/pages/{first["id"]}/recognize').status_code==503
        h=c.get(f'/api/pages/{second["id"]}/history').json();assert len(h['versions'])==4 and len(h['corrections'])==2
    with TestClient(create_app(root)) as c:
        pages=c.get(f'/api/projects/{pid}/pages').json();assert len(pages)==2
        assert pages[1]['status']=='confirmed'
        r=c.get(f'/api/projects/{pid}/export.docx');assert r.status_code==200,r.text
        doc=Document(io.BytesIO(r.content));text='\n'.join(p.text for p in doc.paragraphs)
        assert text.count('Исходная страница 2')==1 and 'Исходная страница 1' not in text
        assert 'температура повышается' in text
        assert len(c.get('/api/profile').json()['examples'])==1
        r2=c.get(f'/api/projects/{pid}/export.docx');doc2=Document(io.BytesIO(r2.content))
        assert '\n'.join(p.text for p in doc2.paragraphs).count('Исходная страница 2')==1


def test_ambiguous_reply_does_not_change_page(client):
    pid=project(client);p=edit(client,upload(client,pid),document([block(),block('b2')],False))
    result=respond(client,p,'Тем — это температура');assert result['needs_clarification']
    assert len(result['block_ids'])==2
    unchanged=client.get(f'/api/pages/{p["id"]}').json();assert unchanged['version']==p['version']
    result=respond(client,p,'Тем — это температура',block_id='b2')['page']
    assert result['document']['blocks'][0]['text']=='тем повышается'
    assert result['document']['blocks'][1]['text']=='температура повышается'


def test_confirm_keeps_literal_x_and_scopes_questions(client):
    pid=project(client);doc=document([block(text='в момент x')]);doc['questions'][0].update(kind='possible_error',original='x',prompt='Написано x, по смыслу ожидается t. Оставить x?')
    p=edit(client,upload(client,pid),doc)
    p=respond(client,p,'оставить как есть',question_id='q1')['page']
    r=client.post(f'/api/pages/{p["id"]}/confirm',json={'version':p['version'],'remember':False})
    assert r.status_code==200 and r.json()['document']['blocks'][0]['text']=='в момент x'
    assert client.get('/api/profile').json()['examples']==[]


def test_conflicting_versions_history_undo_and_original_preserved(client):
    pid=project(client);p=edit(client,upload(client,pid),document(question=False))
    changed=copy.deepcopy(p['document']);changed['blocks'][0]['text']='Исправлено';changed['blocks'][0]['original_text']='подмена исходника'
    new=edit(client,p,changed);assert new['document']['blocks'][0]['original_text']=='тем повышается'
    stale=client.put(f'/api/pages/{p["id"]}',json={'version':p['version'],'document':changed});assert stale.status_code==409
    r=client.post(f'/api/pages/{p["id"]}/restore',json={'version':p['version'],'current_version':new['version']})
    assert r.status_code==200 and r.json()['document']['blocks'][0]['text']=='тем повышается'
    assert r.json()['version']==new['version']+1 and r.json()['status']=='review'
    assert client.get(f'/api/pages/{p["id"]}/history/{p["version"]}').json()['document']['blocks'][0]['text']=='тем повышается'


def test_upload_formats_duplicates_pdf_and_atomic_failure(client):
    pid=project(client);p=upload(client,pid)
    r=client.post(f'/api/projects/{pid}/upload',files={'files':('copy.png',photo(),'image/png')});assert r.json()['duplicates'][0]['exact']
    pdf=fitz.open();pdf.new_page();pdf.new_page()
    r=client.post(f'/api/projects/{pid}/upload',files={'files':('pages.pdf',pdf.tobytes(),'application/pdf')});assert r.status_code==200 and len(r.json()['pages'])==2
    assert r.json()['pages'][0]['original'] if 'original' in r.json()['pages'][0] else True
    heic=photo(format='HEIF');assert client.post(f'/api/projects/{pid}/upload',files={'files':('image.heic',heic,'image/heic')}).status_code==200
    assert client.post(f'/api/projects/{pid}/upload',files={'files':('image.jpg',photo(format='JPEG'),'image/jpeg')}).status_code==200
    before=len(client.get(f'/api/projects/{pid}/pages').json())
    r=client.post(f'/api/projects/{pid}/upload',files=[('files',('ok.png',photo(),'image/png')),('files',('bad.png',b'broken','image/png'))]);assert r.status_code==422
    assert len(client.get(f'/api/projects/{pid}/pages').json())==before
    assert client.get(f'/api/pages/{p["id"]}/original').content==photo()


@pytest.mark.parametrize('latex',FORMULAS)
def test_all_required_formulas_are_editable_omml(latex):
    element=formula_element(latex)
    assert element.tag=='{'+MATH_NS+'}oMath'
    xml=etree.tostring(element).decode()
    if '\\frac' in latex: assert '<m:f>' in xml
    if '\\vec' in latex: assert '<m:acc>' in xml
    if '_x' in latex or '_0' in latex: assert '<m:sSub>' in xml
    if '\\int' in latex: assert '∫' in xml or '&#8747;' in xml


def test_export_formulas_figures_import_and_no_duplicate_pages(client):
    pid=project(client);p=upload(client,pid)
    blocks=[block(f'f{i}','',typ='formula',latex=f,original_latex=f) for i,f in enumerate(FORMULAS)]
    blocks.append(block('fig','График скорости',typ='figure',bbox=[.1,.2,.6,.4],explanation='Площадь — перемещение'))
    p=edit(client,p,document(blocks,False));p=respond(client,p,'Всё верно')['page']
    base=Document();base.add_paragraph('Ранее подтверждённый материал');buf=io.BytesIO();base.save(buf)
    assert client.post(f'/api/projects/{pid}/import-docx',files={'file':('base.docx',buf.getvalue())}).status_code==200
    for _ in range(2):
        r=client.get(f'/api/projects/{pid}/export.docx');assert r.status_code==200
        assert r.headers['X-Formula-Fallbacks']=='0'
        with zipfile.ZipFile(io.BytesIO(r.content)) as z:
            root=etree.fromstring(z.read('word/document.xml'))
            assert len(root.findall('.//{'+MATH_NS+'}oMath'))==7
            assert len([n for n in z.namelist() if n.startswith('word/media/')])==1
        doc=Document(io.BytesIO(r.content));text='\n'.join(x.text for x in doc.paragraphs)
        assert 'Ранее подтверждённый материал' in text and text.count('Исходная страница 1')==1
        assert 'Пояснение приложения' not in text
        assert doc.inline_shapes[0].width<=doc.sections[-1].page_width-doc.sections[-1].left_margin-doc.sections[-1].right_margin
    with_explanations=client.get(f'/api/projects/{pid}/export.docx?explanations=true')
    assert 'Пояснение приложения' in '\n'.join(p.text for p in Document(io.BytesIO(with_explanations.content)).paragraphs)


def test_drafts_and_source_numbers_survive_reorder(client):
    pid=project(client);a=edit(client,upload(client,pid),document());b=edit(client,upload(client,pid,photo(1)),document(question=False))
    assert client.get(f'/api/projects/{pid}/export.docx').status_code==422
    assert client.put(f'/api/projects/{pid}/order',json={'ids':[b['id'],a['id']]}).status_code==200
    r=client.get(f'/api/projects/{pid}/export.docx?drafts=true');assert r.status_code==200
    text='\n'.join(x.text for x in Document(io.BytesIO(r.content)).paragraphs)
    assert text.index('Исходная страница 2')<text.index('Исходная страница 1')
    assert 'ЧЕРНОВИК' in text and '[?]' in text
    assert [p['source_number'] for p in client.get(f'/api/projects/{pid}/pages').json()]==['2','1']


def test_graph_correction_redraw_and_svg_safety(client):
    pid=project(client);diagram={'elements':[{'type':'line','x1':150,'y1':600,'x2':150,'y2':100,'arrow':True},{'type':'text','x':100,'y':90,'text':'Y'}]}
    p=edit(client,upload(client,pid),document([block(typ='figure',text='График',diagram=diagram)],False))
    p=respond(client,p,'На вертикальной оси Y меняем на X')['page']
    assert p['document']['blocks'][0]['diagram']['elements'][1]['text']=='X'
    p['document']['blocks'][0].update(figure_mode='redraw',redraw_confirmed=True)
    p=edit(client,p,p['document']);p=respond(client,p,'Всё верно')['page']
    assert b'X</text>' in client.get(f'/api/pages/{p["id"]}/diagram/b1').content
    assert client.get(f'/api/projects/{pid}/export.docx').status_code==200
    assert client.post('/api/diagram-preview',json={'elements':[{'type':'script','text':'alert(1)'}]}).status_code==422
    for kind in ('position','velocity','parallelogram','cross','components'):
        assert '<svg' in render_svg(template(kind))


def test_backup_restore_keeps_pages_history_profile_and_no_secret(client):
    pid=project(client);p=edit(client,upload(client,pid),document());p=respond(client,p,'Тем — это температура',question_id='q1')['page'];p=respond(client,p,'Всё верно')['page']
    client.put('/api/settings',json={'base_url':'https://api.openai.com/v1','model':'gpt-4.1','api_key':'unit-test-secret','consent':True})
    r=client.get(f'/api/projects/{pid}/backup.zip');assert r.status_code==200
    with zipfile.ZipFile(io.BytesIO(r.content)) as z:
        manifest=z.read('manifest.json');assert b'unit-test-secret' not in manifest and b'encrypted_key' not in manifest
    restored=client.post('/api/backups/restore',files={'file':('backup.zip',r.content)});assert restored.status_code==200,restored.text
    rp=restored.json()['id'];pages=client.get(f'/api/projects/{rp}/pages').json();assert len(pages)==1 and pages[0]['status']=='confirmed'
    assert pages[0]['id']!=p['id']
    h=client.get(f'/api/pages/{pages[0]["id"]}/history').json();assert len(h['versions'])==4 and len(h['corrections'])==2
    assert client.get(f'/api/pages/{pages[0]["id"]}/image').status_code==200
    assert client.get(f'/api/projects/{rp}/export.docx').status_code==200
    assert client.get('/api/profile/export.zip').status_code==200


def test_zip_slip_rejected(client):
    buf=io.BytesIO()
    with zipfile.ZipFile(buf,'w') as z:z.writestr('../escape','bad')
    assert client.post('/api/backups/restore',files={'file':('bad.zip',buf.getvalue())}).status_code==422


def test_provider_keys_encrypted_never_returned_and_consent(client):
    payload={'base_url':'https://api.openai.com/v1','model':'gpt-4.1','api_key':'not-a-real-key','consent':False}
    r=client.put('/api/settings',json=payload);assert r.status_code==200 and r.json()['key_configured']
    assert 'not-a-real-key' not in r.text and 'not-a-real-key' not in client.get('/api/settings').text
    with client.app.state.store.connect() as db: assert 'not-a-real-key' not in db.execute('SELECT value FROM settings').fetchone()[0]
    pid=project(client);p=upload(client,pid)
    r=client.post(f'/api/pages/{p["id"]}/recognize');assert r.status_code==503 and 'Подтвердите отправку' in r.json()['detail']
    payload['base_url']='http://localhost:8000';assert client.put('/api/settings',json=payload).status_code==422
    payload['base_url']='https://127.0.0.1/v1';assert client.put('/api/settings',json=payload).status_code==422


class FixtureProvider(Provider):
    def require(self):return {'key':'test-only'}
    def settings(self,public=False):return {'model':'explicit-test-fixture'}
    async def recognize(self,page,project):return document()


def test_durable_recognition_queue_and_retry_history(tmp_path):
    with TestClient(create_app(tmp_path/'worker',provider_factory=FixtureProvider)) as c:
        pid=project(c);p=upload(c,pid)
        assert c.post(f'/api/pages/{p["id"]}/recognize').status_code==202
        for _ in range(100):
            p=c.get(f'/api/pages/{p["id"]}').json()
            if p['status']!='processing':break
            time.sleep(.03)
        assert p['status']=='review' and p['document']['blocks'][0]['text']=='тем повышается'
        assert len(c.get(f'/api/pages/{p["id"]}/history').json()['recognitions'])==1
        p=respond(c,p,'Тем — это температура',question_id='q1')['page'];v=p['version']
        assert c.post(f'/api/pages/{p["id"]}/recognize').status_code==202
        for _ in range(100):
            p=c.get(f'/api/pages/{p["id"]}').json()
            if p['status']!='processing':break
            time.sleep(.03)
        previous=c.get(f'/api/pages/{p["id"]}/history/{v}').json()
        assert previous['document']['blocks'][0]['text']=='температура повышается'


def test_interrupted_job_requires_manual_retry(tmp_path):
    root=tmp_path/'interrupted';app=create_app(root)
    with TestClient(app) as c:
        pid=project(c);p=upload(c,pid)
        with app.state.store.connect() as db:
            db.execute('INSERT INTO jobs VALUES(?,?,?,?,?,?,?)',(uid(),p['id'],'running',1,now(),now(),None))
            db.execute("UPDATE pages SET status='processing' WHERE id=?",(p['id'],))
    with TestClient(create_app(root)) as c:
        page=c.get(f'/api/pages/{p["id"]}').json();assert page['status']=='error' and 'прервана' in page['error']


def test_password_and_cross_site_protection(tmp_path,monkeypatch):
    monkeypatch.setenv('APP_PASSWORD','test-pass')
    with TestClient(create_app(tmp_path/'auth')) as c:
        assert c.get('/api/projects').status_code==401
        assert c.post('/login',data={'password':'wrong'}).status_code==401
        assert c.post('/login',data={'password':'test-pass'}).status_code==200
        assert c.get('/api/projects').status_code==200
        assert c.post('/api/projects',json={'title':'X'},headers={'Origin':'https://evil.example'}).status_code==403


def test_profile_import_round_trip_is_repeatable(client):
    r=client.post('/api/profile/examples',files={'file':('sample.png',photo())},data={'text':'и u n','scope':'global','context':'Похожие символы'})
    assert r.status_code==201
    ident=r.json()['id']
    assert client.patch('/api/profile/examples/'+ident,json={'text':'и u n v','scope':'subject','subject':'Физика','context':'Скорость v похожа на u'}).status_code==200
    client.post('/api/profile/notes',json={'text':'v похожа на u','scope':'subject','subject':'Физика'})
    backup=client.get('/api/profile/export.zip').content
    path=client.app.state.store.path(client.get('/api/profile').json()['examples'][0]['image'])
    client.delete('/api/profile/examples/'+ident);assert not path.exists()
    for expected in (1,0):
        r=client.post('/api/profile/import',files={'file':('profile.zip',backup)})
        assert r.status_code==200,r.text
        assert r.json()['added']==expected
    data=client.get('/api/profile').json();assert len(data['examples'])==1 and len(data['notes'])==1
    assert data['examples'][0]['text']=='и u n v'


@pytest.mark.parametrize('code,body,message',[
    (429,{},'ограничил'),(401,{},'отклонил'),(500,{},'HTTP 500'),
    (200,{'choices':[{'message':{'content':'not JSON'}}]},'неподдерживаемого'),
    (200,{'choices':[{'finish_reason':'length','message':{'content':'{}'}}]},'обрезан'),
])
def test_provider_errors_are_actionable_and_sanitized(tmp_path,monkeypatch,code,body,message):
    import asyncio,httpx
    from notebook.store import Store
    p=Provider(Store(tmp_path/'p'))
    p.require=lambda:{'key':'synthetic-secret','base_url':'https://unit.example/v1','model':'test'}
    real=httpx.AsyncClient
    def handler(request):
        assert request.headers['Authorization']=='Bearer synthetic-secret'
        return httpx.Response(code,json=body)
    monkeypatch.setattr(httpx,'AsyncClient',lambda **kw:real(transport=httpx.MockTransport(handler),**kw))
    with pytest.raises(ProviderError,match=message) as error:asyncio.run(p.request([{'role':'user','content':'Проверка'}]))
    assert 'synthetic-secret' not in str(error.value)


def test_provider_actual_request_contract_and_personalization(client,monkeypatch):
    import asyncio,httpx
    pid=project(client);page=upload(client,pid)
    client.post('/api/profile/examples',files={'file':('sample.png',photo())},data={'text':'v похожа на u','scope':'subject','subject':'Физика'})
    client.post('/api/profile/notes',json={'text':'x сохранять как x','scope':'global'})
    provider=client.app.state.provider
    provider.require=lambda:{'key':'test-contract','base_url':'https://unit.example/v1','model':'vision-test'}
    real=httpx.AsyncClient;received=[]
    def handler(request):
        payload=json.loads(request.content);received.append(payload)
        assert request.url.path=='/v1/chat/completions'
        return httpx.Response(200,json={'choices':[{'finish_reason':'stop','message':{'content':json.dumps(document(),ensure_ascii=False)}}]})
    monkeypatch.setattr(httpx,'AsyncClient',lambda **kw:real(transport=httpx.MockTransport(handler),**kw))
    result=asyncio.run(provider.recognize(client.app.state.store.page(page['id']),client.app.state.store.project(pid)))
    assert result['blocks'][0]['original_text']=='тем повышается'
    content=received[0]['messages'][1]['content']
    assert len([x for x in content if x['type']=='image_url'])==2  # one confirmed example + current page
    assert content[-1]['image_url']['url'].startswith('data:image/jpeg;base64,')
    assert 'v похожа на u' in json.dumps(content,ensure_ascii=False)
    assert received[0]['response_format']=={'type':'json_object'}


def test_shared_context_notes_seed_once_and_remain_editable(tmp_path):
    root=tmp_path/'seed'
    with TestClient(create_app(root,seed_profile=True)) as c:
        notes=c.get('/api/profile').json()['notes'];assert len(notes)==4
        assert 'm' in ' '.join(n['text'] for n in notes)
        c.delete('/api/profile/notes/'+notes[0]['id'])
        c.patch('/api/profile/notes/'+notes[1]['id'],json={'text':'Моё уточнение'})
    with TestClient(create_app(root,seed_profile=True)) as c:
        notes=c.get('/api/profile').json()['notes'];assert len(notes)==3
        assert any(n['text']=='Моё уточнение' for n in notes)


@pytest.mark.parametrize('status,code,fragment',[
    (401,'invalid_api_key','API-ключ не принят'),
    (403,'unsupported_country_region_territory','страны или региона'),
    (403,'insufficient_permissions','не хватает разрешений'),
    (403,'account_deactivated','отключил аккаунт'),
    (404,'model_not_found','Модель не найдена'),
    (429,'insufficient_quota','квота API'),
    (429,'rate_limit_exceeded','частота запросов'),
])
def test_provider_error_codes_explain_specific_causes_without_echo(status,code,fragment):
    import httpx
    from notebook.provider import explain_http_error
    response=httpx.Response(status,json={'error':{'code':code,'message':'Sensitive key sk-user-secret; personal page text'}})
    result=explain_http_error(response)
    assert f'HTTP {status}' in result and code in result and fragment in result
    assert 'sk-user-secret' not in result and 'personal page text' not in result


@pytest.mark.parametrize('payload',[
    {'error':{'code':'sk-user-secret','message':'Bearer private-value'}},
    {'error':'sk-user-secret'},
    ['sk-user-secret'],
    {'error':{'code':['private-value'],'type':'private-value'}},
])
def test_unknown_error_payload_does_not_expose_provider_text(payload):
    import httpx
    from notebook.provider import explain_http_error
    result=explain_http_error(httpx.Response(403,json=payload))
    assert 'HTTP 403' in result and 'Причина не уточнена' in result
    assert 'sk-user-secret' not in result and 'private-value' not in result


def test_non_json_error_and_known_type_are_safe():
    import httpx
    from notebook.provider import explain_http_error
    assert 'HTTP 403' in explain_http_error(httpx.Response(403,text='<html>private-value</html>'))
    result=explain_http_error(httpx.Response(429,json={'error':{'code':None,'type':'insufficient_quota','message':'private-value'}}))
    assert 'insufficient_quota' in result and 'private-value' not in result


def test_key_whitespace_is_trimmed_and_empty_input_does_not_destroy_binding(client):
    settings={'base_url':'https://api.openai.com/v1','model':'gpt-4.1','api_key':'  synthetic-test-token  ','consent':False}
    r=client.put('/api/settings',json=settings);assert r.status_code==200
    assert client.app.state.provider.settings()['key']=='synthetic-test-token'
    settings['api_key']='   '
    assert client.put('/api/settings',json=settings).status_code==422
    assert client.app.state.provider.settings()['key']=='synthetic-test-token'
