import io
import json
import re
import zipfile
from pathlib import Path, PurePosixPath

import latex2mathml.converter
import mathml2omml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Cm, Inches, Pt
from docxcompose.composer import Composer
from lxml import etree
from PIL import Image

from .diagrams import render_svg, render_png
from .images import crop
from .store import dumps,uid,now

from .math_export import MATH_NS, formula_element


def export_docx(store, project_id, drafts=False, explanations=False):
    project=store.project(project_id)
    with store.connect() as db:
        ids=[r['id'] for r in db.execute('SELECT id FROM pages WHERE project_id=? ORDER BY position,id',(project_id,))]
    pages=[store.page(i) for i in ids]
    pages=[p for p in pages if p['status']=='confirmed' or drafts]
    if not pages and not project.get('base_docx'): raise ValueError('Нет подтверждённых страниц. Подтвердите страницу или включите экспорт черновиков.')
    doc=Document()
    section=doc.sections[0]
    section.top_margin=section.bottom_margin=Cm(2)
    section.left_margin=section.right_margin=Cm(2)
    normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6)
    doc.add_heading(project['title'],0)
    warnings=[]
    for page in pages:
        label=f'Исходная страница {page["source_number"]}'
        heading=doc.add_heading(label,level=2)
        heading.paragraph_format.keep_with_next=True
        if page['status']!='confirmed': doc.add_paragraph('ЧЕРНОВИК — не подтверждён пользователем',style='Subtitle')
        for block in page['document']['blocks']:
            typ=block['type']
            if typ in ('text','heading'):
                if typ=='heading': doc.add_heading(block['text'],level=3)
                else:
                    for line in block['text'].split('\n'): doc.add_paragraph(line)
            elif typ=='formula':
                p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                try: p._p.append(formula_element(block['latex']))
                except Exception:
                    p.add_run('Формула (LaTeX): '+block['latex'])
                    warnings.append(f'Стр. {page["source_number"]}: формула {block["id"]} сохранена как LaTeX')
                if block.get('number'): p.add_run('   ('+block['number']+')')
            elif typ=='figure':
                if block.get('figure_mode')=='redraw' and block.get('redraw_confirmed') and block.get('diagram'):
                    data=render_png(block['diagram'])
                else: data=crop(store,page,block['bbox'])
                with Image.open(io.BytesIO(data)) as img:
                    w,h=img.size
                    max_width=(section.page_width-section.left_margin-section.right_margin)/914400
                    width=min(max_width,5.8,6.8*w/max(h,1))
                p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(io.BytesIO(data),width=Inches(width))
                if block.get('text'): doc.add_paragraph(block['text'],style='Caption')
                if explanations and block.get('explanation'): doc.add_paragraph('Пояснение приложения: '+block['explanation'])
        if page['status']!='confirmed':
            for q in page['document']['questions']:
                if not q['resolved']: doc.add_paragraph('[?] '+q['prompt'])
    if project.get('base_docx'):
        base=Document(store.path(project['base_docx']))
        composer=Composer(base); composer.append(doc); doc=composer.doc
    stream=io.BytesIO(); doc.save(stream)
    return stream.getvalue(),warnings


def backup_project(store, project_id):
    project=store.project(project_id)
    with store.connect() as db:
        pages=[dict(r) for r in db.execute('SELECT * FROM pages WHERE project_id=? ORDER BY position',(project_id,))]
        ids=[p['id'] for p in pages]
        versions=[]; corrections=[]; runs=[]
        for pid in ids:
            versions += [dict(r) for r in db.execute('SELECT * FROM versions WHERE page_id=?',(pid,))]
            corrections += [dict(r) for r in db.execute('SELECT * FROM corrections WHERE page_id=?',(pid,))]
            runs += [dict(r) for r in db.execute('SELECT * FROM recognition_runs WHERE page_id=?',(pid,))]
        examples=[dict(r) for r in db.execute("SELECT * FROM examples WHERE scope='global' OR subject=?",(project['subject'],))]
        notes=[dict(r) for r in db.execute("SELECT * FROM handwriting_notes WHERE scope='global' OR subject=?",(project['subject'],))]
    manifest={'format':'russian-notebook','version':1,'project':project,'pages':pages,'versions':versions,'corrections':corrections,'recognition_runs':runs,'examples':examples,'notes':notes}
    assets={p[k] for p in pages for k in ('original','image')}
    assets.update(e['image'] for e in examples)
    if project.get('base_docx'): assets.add(project['base_docx'])
    stream=io.BytesIO()
    with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('manifest.json',dumps(manifest))
        for name in sorted(assets): z.write(store.path(name),name)
        # Human-readable structured documents are included in addition to exact DB snapshots.
        for p in pages: z.writestr(f'documents/{p["position"]:04d}-{p["id"]}.json',p['document'])
    return stream.getvalue()


def read_zip(data, limit=800*1024*1024):
    try: z=zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e: raise ValueError('Повреждённый ZIP') from e
    infos=z.infolist()
    if len(infos)>12000 or sum(i.file_size for i in infos)>limit:
        z.close(); raise ValueError('Резервная копия слишком велика')
    names=[i.filename for i in infos]
    if len(set(names))!=len(names): z.close(); raise ValueError('Повторяющиеся пути в архиве')
    for i in infos:
        p=PurePosixPath(i.filename)
        if i.flag_bits & 1:
            z.close(); raise ValueError('Архив защищён паролем')
        if p.is_absolute() or '..' in p.parts or '\\' in i.filename or (i.external_attr>>16)&0o170000==0o120000:
            z.close(); raise ValueError('Недопустимый путь в архиве')
    return z


def restore_backup(store,data):
    from .models import PageDocument
    with read_zip(data) as z:
        try: manifest=json.loads(z.read('manifest.json'))
        except (KeyError,ValueError) as e: raise ValueError('Нет корректного manifest.json') from e
        if manifest.get('format')!='russian-notebook' or manifest.get('version')!=1: raise ValueError('Неподдерживаемая версия резервной копии')
        project=manifest['project']; pages=manifest['pages']
        if len(pages)>2000: raise ValueError('Слишком много страниц')
        new_project=uid(); page_ids={p['id']:uid() for p in pages}
        if len(page_ids)!=len(pages): raise ValueError('Повторяющиеся страницы в копии')
        files={}
        def asset(name):
            if name is None: return None
            if name in files: return files[name]
            if not name.startswith('assets/') or name not in z.namelist(): raise ValueError('В копии отсутствует файл')
            suffix=Path(name).suffix
            if suffix not in ('.png','.jpg','.jpeg','.heic','.heif','.pdf','.docx'): raise ValueError('Неизвестный тип вложения')
            files[name]=store.save(z.read(name),suffix)
            return files[name]
        # Validate all page documents and assets before changing database membership.
        for p in pages:
            if p.get('status') not in ('uploaded','processing','review','confirmed','error') or not isinstance(p.get('position'),int) or not isinstance(p.get('version'),int): raise ValueError('Некорректная запись страницы')
            from .images import raster
            raster(z.read(p['image']))
            PageDocument.model_validate(json.loads(p['document']))
            p['original']=asset(p['original']); p['image']=asset(p['image'])
        base=asset(project.get('base_docx'))
        examples=manifest.get('examples',[])
        for e in examples: e['image']=asset(e['image'])
        with store.connect() as db:
            db.execute('INSERT INTO projects VALUES(?,?,?,?,?,?)',(new_project,project['title']+' (восстановлено)',project['subject'],now(),now(),base))
            for p in pages:
                old=p['id']; p['id']=page_ids[old]; p['project_id']=new_project
                p['duplicate_of']=page_ids.get(p.get('duplicate_of'))
                if p['status']=='processing': p['status']='error'; p['error']='Обработка была прервана резервным копированием. Повторите вручную.'
                keys=['id','project_id','position','source_number','filename','original','image','sha','phash','width','height','status','document','version','error','duplicate_of','created','updated']; db.execute(f'INSERT INTO pages ({",".join(keys)}) VALUES ({",".join("?" for _ in keys)})',[p[k] for k in keys])
            for v in manifest.get('versions',[]):
                snap=json.loads(v['snapshot']); snap['id']=page_ids[v['page_id']]; snap['project_id']=new_project
                snap['image']=files.get(snap['image'],snap['image']); snap['original']=files.get(snap['original'],snap['original'])
                db.execute('INSERT INTO versions VALUES(?,?,?,?,?,?)',(uid(),page_ids[v['page_id']],v['version'],dumps(snap),v['action'],v['created']))
            for c in manifest.get('corrections',[]):
                db.execute('INSERT INTO corrections VALUES(?,?,?,?,?,?,?,?)',(uid(),page_ids[c['page_id']],c['block_id'],c['question_id'],c['message'],c['before_text'],c['after_text'],c['created']))
            for r in manifest.get('recognition_runs',[]):
                db.execute('INSERT INTO recognition_runs VALUES(?,?,?,?,?,?)',(uid(),page_ids[r['page_id']],r['document'],r['error'],r['provider'],r['created']))
            for e in examples:
                db.execute('INSERT INTO examples VALUES(?,?,?,?,?,?,?,?,?,?,?)',(uid(),e['subject'],e['scope'],e['text'],e['original_text'],e['context'],e['image'],page_ids.get(e.get('page_id')),e.get('block_id'),e['created'],e['updated']))
            for n in manifest.get('notes',[]):
                if not db.execute('SELECT 1 FROM handwriting_notes WHERE text=? AND scope=? AND subject=?',(n['text'],n['scope'],n['subject'])).fetchone():
                    db.execute('INSERT INTO handwriting_notes VALUES(?,?,?,?,?)',(uid(),n['scope'],n['subject'],n['text'],n['created']))
    return new_project
